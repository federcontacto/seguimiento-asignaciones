from firebase_functions import https_fn, firestore_fn
from firebase_functions.options import set_global_options, SecretParam
from firebase_admin import initialize_app, firestore, auth
import google.oauth2.credentials
import google.auth.transport.requests
import google.auth.exceptions
import google_auth_oauthlib.flow
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload, MediaIoBaseDownload
import urllib.parse
import hashlib
import base64
import copy
import os
import re
import shutil
import tempfile
import unicodedata
import io
import pdfplumber
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import json
import requests as requests_lib
from datetime import datetime, timezone

set_global_options(max_instances=10)
initialize_app()

GOOGLE_CLIENT_ID = SecretParam("GOOGLE_CLIENT_ID")
GOOGLE_CLIENT_SECRET = SecretParam("GOOGLE_CLIENT_SECRET")
SCOPES = ["https://www.googleapis.com/auth/drive"]
APP_URL = "https://siniestros-equipo.web.app"

CALLBACK_URL = "https://driveauthcallback-43i7e7zhfq-uc.a.run.app"


def _flow(client_id: str, client_secret: str):
    return google_auth_oauthlib.flow.Flow.from_client_config(
        {
            "web": {
                "client_id": client_id,
                "client_secret": client_secret,
                "auth_uri": "https://accounts.google.com/o/oauth2/auth",
                "token_uri": "https://oauth2.googleapis.com/token",
                "redirect_uris": [CALLBACK_URL],
            }
        },
        scopes=SCOPES,
        redirect_uri=CALLBACK_URL,
    )


def _pkce_pair():
    """Genera un code_verifier y su code_challenge (S256)."""
    code_verifier = base64.urlsafe_b64encode(os.urandom(40)).rstrip(b"=").decode()
    digest = hashlib.sha256(code_verifier.encode()).digest()
    code_challenge = base64.urlsafe_b64encode(digest).rstrip(b"=").decode()
    return code_verifier, code_challenge


@https_fn.on_request(secrets=[GOOGLE_CLIENT_ID, GOOGLE_CLIENT_SECRET])
def driveAuthStart(req: https_fn.Request) -> https_fn.Response:
    """
    Recibe ?token=<firebase_id_token>, verifica la identidad del verificador,
    y redirige a Google OAuth consent screen.
    """
    id_token = req.args.get("token")
    if not id_token:
        return https_fn.Response("Token requerido", status=400)

    try:
        decoded = auth.verify_id_token(id_token)
        uid = decoded["uid"]
    except Exception:
        return https_fn.Response("Token inválido", status=401)

    code_verifier, code_challenge = _pkce_pair()

    # Guardar code_verifier en Firestore para usarlo en el callback
    db = firestore.client()
    db.collection("oauth_temp").document(uid).set({
        "code_verifier": code_verifier,
    })

    flow = _flow(GOOGLE_CLIENT_ID.value, GOOGLE_CLIENT_SECRET.value)
    auth_url, _ = flow.authorization_url(
        access_type="offline",
        include_granted_scopes="true",
        prompt="consent",
        state=uid,
        code_challenge=code_challenge,
        code_challenge_method="S256",
    )

    return https_fn.Response(
        status=302,
        headers={"Location": auth_url},
        response=""
    )


@https_fn.on_request(secrets=[GOOGLE_CLIENT_ID, GOOGLE_CLIENT_SECRET])
def driveAuthCallback(req: https_fn.Request) -> https_fn.Response:
    """
    Google redirige acá con ?code=...&state=uid.
    Intercambia el code por tokens y guarda el refresh_token en Firestore.
    """
    code = req.args.get("code")
    uid = req.args.get("state")

    if not code or not uid:
        return https_fn.Response("Parámetros faltantes", status=400)

    try:
        db = firestore.client()

        # Recuperar code_verifier guardado en Start
        temp_doc = db.collection("oauth_temp").document(uid).get()
        code_verifier = temp_doc.get("code_verifier") if temp_doc.exists else None
        db.collection("oauth_temp").document(uid).delete()

        flow = _flow(GOOGLE_CLIENT_ID.value, GOOGLE_CLIENT_SECRET.value)
        flow.fetch_token(code=code, code_verifier=code_verifier)
        creds = flow.credentials

        db.collection("usuarios").document(uid).set(
            {
                "drive_token": {
                    "token": creds.token,
                    "refresh_token": creds.refresh_token,
                    "token_uri": creds.token_uri,
                    "scopes": list(creds.scopes) if creds.scopes else SCOPES,
                    # client_id y client_secret NO se persisten: se leen del Secret Manager en runtime
                }
            },
            merge=True,
        )
    except Exception as e:
        print(f"[driveAuthCallback] ERROR: {e}")
        return https_fn.Response("Error interno del servidor", status=500)

    return https_fn.Response(
        status=302,
        headers={"Location": f"{APP_URL}?drive_connected=1"},
        response=""
    )


# ── DRIVE HELPERS ─────────────────────────────────────────────────────────────

DRIVE_ROOT_NAME = "asignaciones-equipo"


def _drive_service_for_uid(uid: str):
    """Devuelve un servicio Drive autenticado con el token del verificador."""
    db = firestore.client()
    doc = db.collection("usuarios").document(uid).get()
    token_data = doc.get("drive_token") if doc.exists else None
    if not token_data or not token_data.get("refresh_token"):
        raise ValueError(f"Verificador {uid} no tiene Drive conectado")

    creds = google.oauth2.credentials.Credentials(
        token=token_data.get("token"),
        refresh_token=token_data["refresh_token"],
        token_uri=token_data.get("token_uri", "https://oauth2.googleapis.com/token"),
        client_id=GOOGLE_CLIENT_ID.value,
        client_secret=GOOGLE_CLIENT_SECRET.value,
        scopes=token_data.get("scopes", SCOPES),
    )
    if creds.expired and creds.refresh_token:
        try:
            creds.refresh(google.auth.transport.requests.Request())
        except google.auth.exceptions.RefreshError:
            db.collection("usuarios").document(uid).set(
                {"drive_token_error": True}, merge=True
            )
            raise ValueError(f"Token de Drive revocado para {uid}. El verificador debe reconectar.")
        db.collection("usuarios").document(uid).set(
            {"drive_token_error": False}, merge=True
        )
        db.collection("usuarios").document(uid).set(
            {"drive_token": {
                "token": creds.token,
                "refresh_token": creds.refresh_token,
                "token_uri": creds.token_uri,
                "scopes": list(creds.scopes) if creds.scopes else SCOPES,
                # client_id y client_secret NO se persisten
            }},
            merge=True,
        )
    return build("drive", "v3", credentials=creds, cache_discovery=False)


def _find_or_create_folder(service, name: str, parent_id: str = None) -> str:
    """Busca una carpeta por nombre (y parent opcional). La crea si no existe."""
    query = (
        f"mimeType='application/vnd.google-apps.folder' "
        f"and name='{name}' and trashed=false"
    )
    if parent_id:
        query += f" and '{parent_id}' in parents"

    results = service.files().list(q=query, fields="files(id)", pageSize=5).execute()
    files = results.get("files", [])
    if files:
        return files[0]["id"]

    metadata = {"name": name, "mimeType": "application/vnd.google-apps.folder"}
    if parent_id:
        metadata["parents"] = [parent_id]
    folder = service.files().create(body=metadata, fields="id").execute()
    return folder["id"]


# ── TRIGGER: CREAR CARPETA AL ASIGNAR CASO ───────────────────────────────────

@firestore_fn.on_document_created(document="casos/{casoId}")
def crearCarpetaCaso(event: firestore_fn.Event) -> None:
    """
    Se dispara solo cuando se crea un caso nuevo.
    Si tiene asignado_a, crea la carpeta en el Drive del verificador
    y guarda el folder_id en el documento.
    """
    after = event.data

    data = after.to_dict()
    uid = data.get("asignado_a")
    if not uid:
        return  # sin asignar, nada que hacer

    if data.get("drive_folder_id"):
        return  # carpeta ya creada

    caso_id = event.params["casoId"]
    id_asignacion = data.get("id") or data.get("idAsignacion") or caso_id
    siniestro = data.get("siniestro") or data.get("nSiniestro") or ""
    apellido = _extraer_apellido(data.get("nombre_completo", ""))
    if siniestro and apellido and apellido != "DESCONOCIDO":
        nombre_carpeta = f"{id_asignacion}.{siniestro}.{apellido}"
    elif siniestro:
        nombre_carpeta = f"{id_asignacion}.{siniestro}"
    else:
        nombre_carpeta = id_asignacion

    try:
        service = _drive_service_for_uid(uid)
        root_id = _find_or_create_folder(service, DRIVE_ROOT_NAME)
        folder_id = _find_or_create_folder(service, nombre_carpeta, parent_id=root_id)

        # Guardar folder_id en el caso
        db = firestore.client()
        db.collection("casos").document(caso_id).update({
            "drive_folder_id": folder_id,
        })
    except Exception as e:
        print(f"[crearCarpetaCaso] Error para caso {caso_id}: {e}")


_CORS_CARPETA = {
    "Access-Control-Allow-Origin": "https://siniestros-equipo.web.app",
    "Access-Control-Allow-Methods": "POST, OPTIONS",
    "Access-Control-Allow-Headers": "Content-Type, Authorization",
}

@https_fn.on_request(secrets=[GOOGLE_CLIENT_ID, GOOGLE_CLIENT_SECRET])
def crearCarpetaManual(req: https_fn.Request) -> https_fn.Response:
    """Crea (o encuentra) la carpeta Drive de un caso, invocado manualmente desde el frontend."""
    if req.method == 'OPTIONS':
        return https_fn.Response("", status=204, headers=_CORS_CARPETA)
    if req.method != 'POST':
        return https_fn.Response('Método no permitido', status=405, headers=_CORS_CARPETA)

    # Verificar Firebase ID token
    try:
        id_token = req.headers.get('Authorization', '').replace('Bearer ', '')
        decoded = auth.verify_id_token(id_token)
        uid = decoded['uid']
    except Exception:
        return https_fn.Response(
            json.dumps({'error': 'No autorizado'}), status=401,
            content_type='application/json', headers=_CORS_CARPETA)

    try:
        body = req.get_json(silent=True) or {}
        caso_id = body.get('casoId', '').strip()
        if not caso_id:
            return https_fn.Response(
                json.dumps({'error': 'casoId requerido'}), status=400,
                content_type='application/json', headers=_CORS_CARPETA)

        db_client = firestore.client()
        caso_ref = db_client.collection('casos').document(caso_id)
        caso_snap = caso_ref.get()
        if not caso_snap.exists:
            return https_fn.Response(
                json.dumps({'error': 'Caso no encontrado'}), status=404,
                content_type='application/json', headers=_CORS_CARPETA)

        data = caso_snap.to_dict()
        id_asignacion = data.get('id') or data.get('idAsignacion') or caso_id
        siniestro = data.get('siniestro') or data.get('nSiniestro') or ''
        apellido = _extraer_apellido(data.get('nombre_completo', ''))
        if siniestro and apellido and apellido != 'DESCONOCIDO':
            nombre_carpeta = f"{id_asignacion}.{siniestro}.{apellido}"
        elif siniestro:
            nombre_carpeta = f"{id_asignacion}.{siniestro}"
        else:
            nombre_carpeta = id_asignacion

        # Usar el Drive del verificador asignado al caso (consistente con procesarPDF).
        # Si el caso no tiene asignado_a, usar el Drive del caller como fallback.
        uid_drive = data.get('asignado_a') or uid
        service   = _drive_service_for_uid(uid_drive)
        root_id   = _find_or_create_folder(service, DRIVE_ROOT_NAME)
        folder_id = _find_or_create_folder(service, nombre_carpeta, parent_id=root_id)

        caso_ref.update({'drive_folder_id': folder_id})

        return https_fn.Response(
            json.dumps({'ok': True}), status=200,
            content_type='application/json', headers=_CORS_CARPETA)

    except ValueError as e:
        # _drive_service_for_uid lanza ValueError si Drive no está conectado
        return https_fn.Response(
            json.dumps({'error': str(e)}), status=403,
            content_type='application/json', headers=_CORS_CARPETA)
    except Exception as e:
        print(f"[crearCarpetaManual] Error: {e}")
        return https_fn.Response(
            json.dumps({'error': 'Error interno'}), status=500,
            content_type='application/json', headers=_CORS_CARPETA)


# ── PDF PROCESSING ────────────────────────────────────────────────────────────

PLANTILLAS_DIR = os.path.join(os.path.dirname(__file__), "plantillas")
PLANTILLA_POR_DEFECTO = "choque"
PLANTILLAS = {
    "choque":       "choque.docx",
    "granizo":      "granizo.docx",
    "incendio":     "incendio.docx",
    "robo_total":   "robo_total.docx",
    "robo_parcial": "robo_parcial.docx",
}


def _sin_acentos(txt: str) -> str:
    return unicodedata.normalize("NFKD", txt or "").encode("ascii", "ignore").decode()


def _clave_plantilla(tipo: str) -> str:
    """
    Tipo del caso (config/tarifas: 'Robo Total', 'Robo Parcial', ...) -> plantilla.
    Prosecuciones y ampliaciones siguen con la de choque, como hasta ahora.
    """
    t = _sin_acentos(tipo).lower()
    if not t or "prosecucion" in t or "ampliacion" in t:
        return PLANTILLA_POR_DEFECTO
    if "granizo" in t:
        return "granizo"
    if "parcial" in t or "rueda" in t:
        return "robo_parcial"
    if "incendio" in t:
        return "incendio"
    if "robo" in t:
        return "robo_total"
    return PLANTILLA_POR_DEFECTO


def _plantilla_path(clave: str) -> str:
    return os.path.join(PLANTILLAS_DIR, PLANTILLAS.get(clave, PLANTILLAS[PLANTILLA_POR_DEFECTO]))

MESES = {
    '01':'enero','02':'febrero','03':'marzo','04':'abril','05':'mayo','06':'junio',
    '07':'julio','08':'agosto','09':'septiembre','10':'octubre','11':'noviembre','12':'diciembre'
}


def _normalizar_patente(raw: str) -> str:
    """
    Las patentes viejas (AAA123) llegan en el PDF con el numero rellenado con
    ceros a la izquierda (AAA000123). Sacamos el relleno y dejamos el valor real.
    Formatos AR: vieja AAA123 (3 letras + 3 nros), nueva AA123BC (2 letras + 3 nros + 2 letras).
    """
    if not raw:
        return ''
    raw = raw.strip().upper().replace(' ', '')
    m = re.match(r'^([A-Z]{2,3})(\d+)([A-Z]{0,2})$', raw)
    if not m:
        return raw
    pre, nums, suf = m.group(1), m.group(2), m.group(3)
    nums = nums.lstrip('0')
    if len(nums) < 3:
        nums = nums.zfill(3)  # conservar ceros significativos: 045, 007
    return f"{pre}{nums}{suf}"


_MENORES = {'de', 'del', 'la', 'las', 'los', 'y', 'e', 'en'}


def _titulo(texto: str) -> str:
    """MAYUSCULAS -> Título, dejando partículas en minúscula ('Ciudad de la Paz')."""
    salida = []
    for i, w in enumerate((texto or '').split()):
        wl = w.lower()
        salida.append(wl if (i and wl in _MENORES) else wl.capitalize())
    return ' '.join(salida)


def _formatear_nombre(nombre: str) -> str:
    """'GOMEZ, MARIA LAURA' -> 'GOMEZ, Maria Laura' (apellido en mayúsculas)."""
    n = (nombre or '').strip()
    if not n:
        return ''
    if ',' in n:
        ape, nom = n.split(',', 1)
        return f"{ape.strip().upper()}, {_titulo(nom)}"
    partes = n.split()
    if len(partes) < 2:
        return n.upper()
    # Sin coma el apellido es la última palabra (mismo criterio que _extraer_apellido)
    return f"{_titulo(' '.join(partes[:-1]))} {partes[-1].upper()}"


def _agregar_nro(calle: str) -> str:
    """'Av. Siempreviva 742' -> 'Av. Siempreviva N°742'."""
    if not calle or re.search(r'N[°ºro\.]{1,3}\s*\d', calle, re.IGNORECASE):
        return calle
    return re.sub(r'(?<![\d°º])(\d+)', r'N°\1', calle, count=1)


def _formatear_domicilio(calle: str, localidad: str, provincia: str) -> str:
    """Arma el domicilio con capitalización, N° de altura y provincia expandida."""
    prov_raw = (provincia or '').strip().upper().rstrip('.')
    prov = 'Provincia de Buenos Aires' if prov_raw == 'BUENOS AIRES' else _titulo(provincia)
    return f"{_agregar_nro(_titulo(calle))}, {_titulo(localidad)}, {prov}"


def _formatear_dni(dni: str) -> str:
    """12345678 -> 12.345.678 (separador de miles con punto, formato AR)."""
    d = re.sub(r'\D', '', dni or '')
    if not d:
        return dni or ''
    return f"{int(d):,}".replace(',', '.')


def _extraer_datos_denuncia(pdf_path: str) -> dict:
    try:
        with pdfplumber.open(pdf_path) as pdf:
            texto = "\n".join(p.extract_text() or "" for p in pdf.pages)
    except Exception as e:
        print(f"[pdf] Error leyendo PDF: {e}")
        return {}

    datos = {}
    m = re.search(r'P[OÓ]LIZA\s+([\d-]+)', texto)
    if m: datos['poliza'] = m.group(1)

    m = re.search(r'FECHA SINIESTRO\s+(\d{2}/\d{2}/\d{4})\s+HORA\s+(\d{2}:\d{2})', texto)
    if m:
        datos['fecha_ocurrencia'] = m.group(1)
        datos['hora_ocurrencia']  = m.group(2)

    m = re.search(r'FECHA DENUNCIA\s+(\d{2}/\d{2}/\d{4})', texto)
    if m: datos['fecha_denuncia'] = m.group(1)

    m = re.search(r'MARCA\s+\d+\s+-\s+(.+?)\s+MODELO\s+(\d{4})', texto)
    if m: datos['marca_modelo'] = f"{m.group(1).strip()} {m.group(2)}"

    m = re.search(r'PATENTE\s+([A-Z]{2,3}\d+[A-Z]{0,2})', texto)
    if m: datos['patente'] = _normalizar_patente(m.group(1))

    partes = re.split(r'D\s+A\s+T\s+O\s+S\s+D\s+E\s+L\s+A\s+S\s+E\s+G\s+U\s+R\s+A\s+D\s+O', texto)
    if len(partes) >= 2:
        sec = partes[1]
        m = re.search(r'APELLIDO Y NOMBRE\s+(.+?)(?:\n|DOMICILIO)', sec)
        if m: datos['nombre_completo'] = _formatear_nombre(m.group(1).strip())
        m = re.search(r'NRO\.\s*DOC\.\s+(\d+)', sec)
        if m: datos['dni'] = _formatear_dni(m.group(1))
        m = re.search(r'DOMICILIO\s+(.+?)\s+LOCALIDAD\s+(.+?)\s+PROVINCIA\s+(.+?)(?:\n|TIPO)', sec)
        if m: datos['domicilio'] = _formatear_domicilio(m.group(1).strip(), m.group(2).strip(), m.group(3).strip())
        m = re.search(r'DOM\.ELECTR[OÓ]NICO\s+(\S+@\S+)', sec)
        if m: datos['mail'] = m.group(1)
        sexo = None
        m = re.search(r'\bSEXO\s+(MASCULINO|FEMENINO)\b', sec, re.IGNORECASE)
        if m: sexo = 'M' if m.group(1).upper().startswith('MASC') else 'F'
        if not sexo:
            m = re.search(r'\bSEXO\s+([MF])\b', sec)
            if m: sexo = m.group(1).upper()
        if sexo:
            datos['sexo'] = sexo
            datos['tratamiento'] = 'Sr.' if sexo == 'M' else 'Sra.'

    m = re.search(r'FORMA OCURRENCIA\s+(.+?)\s*(?=OBSERVACIONES)', texto, re.DOTALL)
    if m: datos['relato'] = ' '.join(m.group(1).strip().split())

    return datos


def _extraer_apellido(nombre_completo: str) -> str:
    if not nombre_completo or not nombre_completo.strip():
        return 'DESCONOCIDO'
    if ',' in nombre_completo:
        apellido = nombre_completo.split(',')[0].strip()
    else:
        apellido = nombre_completo.strip().split()[-1]
    return re.sub(r'[^\w]', '', apellido, flags=re.UNICODE).upper() or 'DESCONOCIDO'


def _set_cell(cell, texto: str):
    from docx.oxml.ns import qn
    tc = cell._tc
    parrafos = tc.findall(qn('w:p'))
    for p in parrafos[1:]:
        tc.remove(p)
    para = cell.paragraphs[0]
    for run in para.runs:
        run.text = ""
    if para.runs:
        para.runs[0].text = texto
    else:
        para.add_run(texto)


def _set_paragraph_text(para, texto: str):
    """Reescribe el texto de un parrafo conservando el primer run (formato/alineacion)."""
    for run in para.runs:
        run.text = ""
    if para.runs:
        para.runs[0].text = texto
    else:
        para.add_run(texto)


def _iter_paragraphs(container):
    """Recorre todos los parrafos del cuerpo y de las tablas (incluye tablas anidadas)."""
    for p in container.paragraphs:
        yield p
    for t in container.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _actualizar_fecha_carta(doc):
    """
    La plantilla tiene la fecha escrita a mano ('... abril de 2026').
    La reemplaza por el mes/año actual, conservando el resto de la línea.
    """
    from datetime import date
    hoy = date.today()
    nuevo = f"{MESES[f'{hoy.month:02d}']} de {hoy.year}"
    patron = re.compile(r'(' + '|'.join(MESES.values()) + r')\s+de\s+\d{4}', re.IGNORECASE)
    for p in doc.paragraphs:
        if 'Buenos Aires' in p.text and patron.search(p.text):
            _set_paragraph_text(p, patron.sub(nuevo, p.text))
            return


CONCLUSIONES_TITULO = 'DILIGENCIAS COMPLEMENTARIAS'
CONCLUSIONES_FIN    = 'OPINI'


def _conclusiones_por_defecto(db, clave: str = PLANTILLA_POR_DEFECTO) -> list:
    """
    Textos de las conclusiones marcadas con la estrella en el módulo web
    (config/conclusiones → items[].porDefecto). Si falla, devuelve []
    y el DOCX conserva el texto modelo de la plantilla.

    Una conclusión puede traer `tipo` (ej. 'Incendio'): si hay favoritas del
    tipo del caso se usan esas; si no, las favoritas sin tipo (las de siempre).
    """
    try:
        doc = db.collection("config").document("conclusiones").get()
        if not doc.exists:
            return []
        items = (doc.to_dict() or {}).get("items", []) or []
        favoritas = [i for i in items
                     if i.get("porDefecto") and (i.get("texto") or "").strip()]
        del_tipo = [i for i in favoritas
                    if i.get("tipo") and _clave_plantilla(i["tipo"]) == clave]
        elegidas = del_tipo or [i for i in favoritas if not i.get("tipo")]
        return [i["texto"].strip() for i in elegidas]
    except Exception as e:
        print(f"[conclusiones] no se pudieron leer: {e}")
        return []


def _norm(txt: str) -> str:
    """Normaliza para comparar títulos espaciados tipo 'O P I N I Ó N'."""
    return re.sub(r'\s+', '', (txt or '')).upper()


_VINETA_RE = re.compile(r'^[\s ]*(?:[❖•‣▪●·*]+|[-–—]+(?=\s))[\s ]*')


def _limpiar_vineta(linea: str) -> str:
    """Saca la viñeta tipeada al inicio de la línea; Word pone la suya."""
    return _VINETA_RE.sub('', (linea or '').strip()).strip()


def _reemplazar_conclusiones(doc, textos: list):
    """
    Reemplaza los párrafos modelo entre 'DILIGENCIAS COMPLEMENTARIAS y CONCLUSIONES'
    y 'OPINIÓN' por las conclusiones marcadas por defecto en la web.
    Sin textos no toca nada (la plantilla queda como está).
    """
    if not textos:
        return

    parrafos = doc.paragraphs
    ini = fin = None
    for i, p in enumerate(parrafos):
        n = _norm(p.text)
        if ini is None and _norm(CONCLUSIONES_TITULO) in n:
            ini = i
        elif ini is not None and n.startswith(_norm(CONCLUSIONES_FIN)):
            fin = i
            break
    if ini is None or fin is None or fin <= ini:
        print("[conclusiones] no se ubicó la sección en la plantilla, se deja intacta")
        return

    # Una línea del texto = una viñeta. La ❖ la pone Word (los párrafos de la
    # plantilla son lista real, numId=6), así que se saca la que venga tipeada
    # en la conclusión: si no, sale doble.
    lineas = [_limpiar_vineta(l) for t in textos for l in t.splitlines() if l.strip()]
    lineas = [l for l in lineas if l]
    if not lineas:
        return

    # Párrafos con contenido dentro del bloque: son los que se reemplazan
    bloque = [p for p in parrafos[ini + 1:fin] if p.text.strip()]
    if not bloque:
        print("[conclusiones] bloque modelo vacío, se deja intacto")
        return

    modelo = bloque[0]
    for i, linea in enumerate(lineas):
        if i < len(bloque):
            _set_paragraph_text(bloque[i], linea)
        else:
            # Más líneas que párrafos: clonar el primero para heredar formato
            nuevo = copy.deepcopy(modelo._p)
            bloque[-1]._p.addnext(nuevo)
            from docx.text.paragraph import Paragraph
            para = Paragraph(nuevo, modelo._parent)
            _set_paragraph_text(para, linea)
            bloque.append(para)

    # Menos líneas que párrafos: eliminar los sobrantes
    for p in bloque[len(lineas):]:
        p._p.getparent().remove(p._p)


def _set_resultado_box(doc, resultado: str):
    """Caja 1x1 que en la plantilla dice 'SIN FRAUDE'. Solo se toca si el caso ya tiene resultado."""
    if not resultado:
        return
    for t in doc.tables:
        if len(t.rows) == 1 and len(t.columns) == 1:
            _set_cell(t.rows[0].cells[0], resultado.upper())
            return


def _completar_docx(dest_path: str, datos: dict, case_id: str, siniestro: str,
                    conclusiones: list = None, resultado: str = '',
                    plantilla: str = PLANTILLA_POR_DEFECTO):
    shutil.copy2(_plantilla_path(plantilla), dest_path)
    doc = Document(dest_path)
    tables = doc.tables
    if len(tables) > 0:
        t = tables[0]
        _set_cell(t.rows[0].cells[1], siniestro)
        _set_cell(t.rows[1].cells[1], datos.get('poliza', ''))
        _set_cell(t.rows[2].cells[1], case_id)
        _set_cell(t.rows[3].cells[1], datos.get('fecha_ocurrencia', ''))
        _set_cell(t.rows[4].cells[1], datos.get('fecha_denuncia', ''))
        # 'COBERTURA Y USO' la completa el verificador: se vacía el valor modelo
        if len(t.rows) > 5 and _norm(t.rows[5].cells[0].text).startswith('COBERTURA'):
            _set_cell(t.rows[5].cells[1], datos.get('cobertura', ''))
    if len(tables) > 1:
        t = tables[1]
        _set_cell(t.rows[0].cells[1], datos.get('nombre_completo', ''))
        _set_cell(t.rows[1].cells[1], datos.get('dni', ''))
        _set_cell(t.rows[2].cells[1], datos.get('domicilio', ''))
        _set_cell(t.rows[4].cells[1], datos.get('marca_modelo', ''))
        _set_cell(t.rows[5].cells[1], datos.get('patente', ''))
    if len(tables) > 2 and datos.get('relato'):
        _set_cell(tables[2].rows[2].cells[1], datos['relato'])
    _reemplazar_conclusiones(doc, conclusiones or [])
    _set_resultado_box(doc, resultado)
    _actualizar_fecha_carta(doc)
    # Sin retoques de fuentes ni anchos: la plantilla de cada tipo ya viene
    # con el formato final y lo insertado hereda el formato de su celda/párrafo.
    doc.save(dest_path)


def _descargar_archivo_drive(service, file_id: str, dest_path: str):
    request = service.files().get_media(fileId=file_id)
    with open(dest_path, "wb") as f:
        downloader = MediaIoBaseDownload(f, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()


# ── CLOUD FUNCTION: PROCESAR PDF ─────────────────────────────────────────────

CORS_HEADERS = {
    "Access-Control-Allow-Origin": "https://siniestros-equipo.web.app",
    "Access-Control-Allow-Methods": "POST, OPTIONS",
    "Access-Control-Allow-Headers": "Content-Type, Authorization",
}

@https_fn.on_request(secrets=[GOOGLE_CLIENT_ID, GOOGLE_CLIENT_SECRET])
def procesarPDF(req: https_fn.Request) -> https_fn.Response:
    """
    Recibe {casoId, idToken} por POST.
    Busca DENUNCIA.pdf en la carpeta Drive del caso, extrae datos,
    crea el DOCX, lo sube a Drive y actualiza Firestore.
    """
    if req.method == "OPTIONS":
        return https_fn.Response("", status=204, headers=CORS_HEADERS)

    data = req.get_json(silent=True) or {}
    caso_id  = data.get("casoId")
    id_token = data.get("idToken")
    force    = bool(data.get("force"))

    if not caso_id or not id_token:
        return https_fn.Response('{"error":"casoId e idToken requeridos"}', status=400, content_type="application/json", headers=CORS_HEADERS)

    try:
        decoded = auth.verify_id_token(id_token)
        caller_uid = decoded['uid']
    except Exception:
        return https_fn.Response('{"error":"Token invalido"}', status=401, content_type="application/json", headers=CORS_HEADERS)

    db = firestore.client()
    caso_doc = db.collection("casos").document(caso_id).get()
    if not caso_doc.exists:
        return https_fn.Response('{"error":"Caso no encontrado"}', status=404, content_type="application/json", headers=CORS_HEADERS)

    caso = caso_doc.to_dict()
    uid         = caso.get("asignado_a")

    # Verificar que el caller sea el dueño del caso o supervisor
    caller_doc = db.collection("usuarios").document(caller_uid).get()
    caller_rol = caller_doc.to_dict().get("rol") if caller_doc.exists else None
    if caller_uid != uid and caller_rol != "supervisor":
        return https_fn.Response('{"error":"Sin permiso sobre este caso"}', status=403, content_type="application/json", headers=CORS_HEADERS)
    folder_id   = caso.get("drive_folder_id")
    siniestro   = caso.get("siniestro") or caso.get("nSiniestro") or ""

    if not uid:
        return https_fn.Response('{"error":"Caso sin verificador asignado"}', status=400, content_type="application/json", headers=CORS_HEADERS)

    try:
        print(f"[procesarPDF] caso={caso_id} uid={uid} folder={folder_id}")
        service = _drive_service_for_uid(uid)

        if not folder_id:
            # No hay carpeta registrada: buscar o crear en el Drive del verificador
            nombre_base = f"{caso_id}.{siniestro}" if siniestro else caso_id
            root_id   = _find_or_create_folder(service, DRIVE_ROOT_NAME)
            folder_id = _find_or_create_folder(service, nombre_base, parent_id=root_id)
            db.collection("casos").document(caso_id).update({"drive_folder_id": folder_id})
            print(f"[procesarPDF] carpeta creada/encontrada: {folder_id}")

        # Buscar DENUNCIA.pdf en la carpeta
        query = f"'{folder_id}' in parents and trashed=false"
        archivos = service.files().list(q=query, fields="files(id,name)", pageSize=50).execute().get("files", [])
        print(f"[procesarPDF] archivos en carpeta: {[a['name'] for a in archivos]}")
        denuncia = next((a for a in archivos if "DENUNCIA" in a["name"].upper() and a["name"].upper().endswith(".PDF")), None)

        if not denuncia:
            return https_fn.Response('{"error":"No se encontró DENUNCIA.pdf en la carpeta"}', status=404, content_type="application/json", headers=CORS_HEADERS)

        tmp_dir = tempfile.mkdtemp()
        denuncia_local = os.path.join(tmp_dir, "denuncia.pdf")
        _descargar_archivo_drive(service, denuncia["id"], denuncia_local)

        datos = _extraer_datos_denuncia(denuncia_local)
        nombre_completo = datos.get("nombre_completo", "")
        apellido = _extraer_apellido(nombre_completo)

        # Verificar si ya existe DOCX en la carpeta.
        # Sin force: no pisar el informe que el verificador pudo haber redactado.
        docx_existente = next((a for a in archivos if a["name"].upper().endswith(".DOCX")), None)
        if docx_existente and not force:
            return https_fn.Response(
                json.dumps({
                    "error": f"Ya existe un DOCX en la carpeta ({docx_existente['name']}). Reprocesar lo reemplaza.",
                    "code": "DOCX_EXISTS",
                }),
                status=409, content_type="application/json", headers=CORS_HEADERS)
        if docx_existente:
            service.files().delete(fileId=docx_existente["id"]).execute()

        nombre_docx = f"{caso_id}.{siniestro}.{apellido}.docx"
        docx_local = os.path.join(tmp_dir, nombre_docx)
        plantilla = _clave_plantilla(caso.get("tipo") or "")
        print(f"[procesarPDF] tipo={caso.get('tipo')!r} -> plantilla {plantilla}")
        _completar_docx(docx_local, datos, caso_id, siniestro,
                        conclusiones=_conclusiones_por_defecto(db, plantilla),
                        resultado=caso.get("resultado") or "",
                        plantilla=plantilla)

        media = MediaFileUpload(docx_local, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document", resumable=True)
        service.files().create(body={"name": nombre_docx, "parents": [folder_id]}, media_body=media, fields="id").execute()

        # Renombrar carpeta si tiene apellido nuevo
        id_asignacion = caso.get("id") or caso.get("idAsignacion") or caso_id
        nuevo_nombre_carpeta = f"{id_asignacion}.{siniestro}.{apellido}"
        service.files().update(fileId=folder_id, body={"name": nuevo_nombre_carpeta}).execute()

        # Actualizar Firestore
        from datetime import date
        updates = {
            "nombre_completo":       nombre_completo,
            "mail":                  datos.get("mail", ""),
            "dni":                   datos.get("dni", ""),
            "domicilio":             datos.get("domicilio", ""),
            "marca_modelo":          datos.get("marca_modelo", ""),
            "patente":               datos.get("patente", ""),
            "poliza":                datos.get("poliza", ""),
            "fecha_ocurrencia":      datos.get("fecha_ocurrencia", ""),
            "hora_ocurrencia":       datos.get("hora_ocurrencia", ""),
            "fecha_denuncia":        datos.get("fecha_denuncia", ""),
            "sexo":                  datos.get("sexo", ""),
            "tratamiento":           datos.get("tratamiento", ""),
        }
        db.collection("casos").document(caso_id).update(updates)

        shutil.rmtree(tmp_dir, ignore_errors=True)

        return https_fn.Response(
            json.dumps({"ok": True, "nombre_completo": nombre_completo, "docx": nombre_docx,
                        "plantilla": plantilla}),
            status=200,
            content_type="application/json",
            headers=CORS_HEADERS,
        )

    except Exception as e:
        import traceback
        print(f"[procesarPDF] ERROR: {e}")
        print(traceback.format_exc())
        return https_fn.Response(json.dumps({"error": "Error interno del servidor"}), status=500, content_type="application/json", headers=CORS_HEADERS)


# ── CLOUD FUNCTION: DOCUMENTACIÓN TÉCNICA (solo usuarios activos) ────────────

DOCS_TECNICA_PATH = os.path.join(os.path.dirname(__file__), "docs_tecnica.html")

_CORS_DOCS = {
    "Access-Control-Allow-Origin": APP_URL,
    "Access-Control-Allow-Methods": "GET, OPTIONS",
    "Access-Control-Allow-Headers": "Authorization",
}

@https_fn.on_request()
def docsTecnica(req: https_fn.Request) -> https_fn.Response:
    """Sirve la documentación técnica solo a usuarios autenticados y activos.
    El HTML vive en functions/ (no en Hosting), así que sin token no hay acceso."""
    if req.method == "OPTIONS":
        return https_fn.Response("", status=204, headers=_CORS_DOCS)
    if req.method != "GET":
        return https_fn.Response("Método no permitido", status=405, headers=_CORS_DOCS)

    try:
        id_token = req.headers.get("Authorization", "").replace("Bearer ", "")
        decoded = auth.verify_id_token(id_token)
    except Exception:
        return https_fn.Response('{"error":"No autorizado"}', status=401,
                                 content_type="application/json", headers=_CORS_DOCS)

    # Claims en el JWT; fallback al doc de Firestore para sesiones sin claims
    rol = decoded.get("rol")
    activo = decoded.get("activo")
    if rol is None or activo is None:
        user_doc = firestore.client().collection("usuarios").document(decoded["uid"]).get()
        data = user_doc.to_dict() if user_doc.exists else {}
        rol = data.get("rol")
        activo = data.get("activo")
    if rol not in ("supervisor", "verificador") or not activo:
        return https_fn.Response('{"error":"Sin acceso"}', status=403,
                                 content_type="application/json", headers=_CORS_DOCS)

    try:
        with open(DOCS_TECNICA_PATH, encoding="utf-8") as f:
            html = f.read()
    except Exception as e:
        print(f"[docsTecnica] ERROR leyendo archivo: {e}")
        return https_fn.Response("Error interno", status=500, headers=_CORS_DOCS)

    return https_fn.Response(html, status=200, content_type="text/html; charset=utf-8",
                             headers={**_CORS_DOCS, "Cache-Control": "private, no-store"})


# ── CLOUD FUNCTION: CREAR VERIFICADOR ───────────────────────────────────────

@https_fn.on_request()
def crearVerificador(req: https_fn.Request) -> https_fn.Response:
    """
    Crea un usuario en Firebase Auth + doc en Firestore.
    Solo puede ser invocado por supervisores activos.
    Recibe POST { idToken, nombre, email, password, rol }
    """
    cors = {
        "Access-Control-Allow-Origin": APP_URL,
        "Access-Control-Allow-Methods": "POST, OPTIONS",
        "Access-Control-Allow-Headers": "Content-Type, Authorization",
    }

    if req.method == "OPTIONS":
        return https_fn.Response("", status=204, headers=cors)

    def err(msg, code=400):
        return https_fn.Response(json.dumps({"error": msg}), status=code,
                                 content_type="application/json", headers=cors)

    data = req.get_json(silent=True) or {}
    id_token = data.get("idToken", "")
    if not id_token:
        return err("Token requerido", 401)

    try:
        decoded = auth.verify_id_token(id_token)
        caller_uid = decoded["uid"]
    except Exception:
        return err("Token inválido", 401)

    db = firestore.client()
    caller_doc = db.collection("usuarios").document(caller_uid).get()
    caller_data = caller_doc.to_dict() if caller_doc.exists else {}
    if caller_data.get("rol") != "supervisor" or not caller_data.get("activo"):
        return err("Solo supervisores activos pueden crear usuarios", 403)

    nombre   = data.get("nombre", "").strip()
    email    = data.get("email", "").strip()
    password = data.get("password", "")
    rol      = data.get("rol", "verificador")

    if not nombre or not email or not password:
        return err("nombre, email y contraseña son obligatorios")
    if len(password) < 6:
        return err("La contraseña debe tener al menos 6 caracteres")

    try:
        user_record = auth.create_user(email=email, password=password)
        new_uid = user_record.uid
        db.collection("usuarios").document(new_uid).set({
            "uid": new_uid, "nombre": nombre, "email": email,
            "rol": rol, "activo": True,
        })
        # Setear custom claims para que las Firestore Rules no necesiten hacer get()
        auth.set_custom_user_claims(new_uid, {"rol": rol, "activo": True})
        return https_fn.Response(json.dumps({"ok": True, "uid": new_uid}),
                                 status=200, content_type="application/json", headers=cors)
    except auth.EmailAlreadyExistsError:
        return err("Ese email ya tiene una cuenta registrada", 409)
    except Exception as e:
        print(f"[crearVerificador] ERROR: {e}")
        return err(f"Error interno: {e}", 500)


# ── CLOUD FUNCTION: SET CUSTOM CLAIMS ────────────────────────────────────────

_CORS_CLAIMS = {
    "Access-Control-Allow-Origin": APP_URL,
    "Access-Control-Allow-Methods": "POST, OPTIONS",
    "Access-Control-Allow-Headers": "Content-Type, Authorization",
}

@https_fn.on_request()
def setUserClaims(req: https_fn.Request) -> https_fn.Response:
    """
    Setea custom claims (rol, activo) en el token JWT de un usuario.
    Solo puede ser invocado por supervisores activos.
    Recibe POST { idToken, uid }
    Los claims se leen del documento usuarios/{uid} en Firestore.
    """
    if req.method == "OPTIONS":
        return https_fn.Response("", status=204, headers=_CORS_CLAIMS)

    def err(msg, code=400):
        return https_fn.Response(json.dumps({"error": msg}), status=code,
                                 content_type="application/json", headers=_CORS_CLAIMS)

    data = req.get_json(silent=True) or {}
    id_token = data.get("idToken", "")
    target_uid = data.get("uid", "").strip()

    if not id_token:
        return err("idToken requerido", 401)
    if not target_uid:
        return err("uid requerido")

    try:
        decoded = auth.verify_id_token(id_token)
        caller_uid = decoded["uid"]
    except Exception:
        return err("Token inválido", 401)

    db = firestore.client()
    caller_doc = db.collection("usuarios").document(caller_uid).get()
    caller_data = caller_doc.to_dict() if caller_doc.exists else {}
    if caller_data.get("rol") != "supervisor" or not caller_data.get("activo"):
        return err("Solo supervisores activos pueden setear claims", 403)

    target_doc = db.collection("usuarios").document(target_uid).get()
    if not target_doc.exists:
        return err("Usuario no encontrado", 404)

    target_data = target_doc.to_dict()
    claims = {
        "rol": target_data.get("rol", "verificador"),
        "activo": bool(target_data.get("activo", True)),
    }

    try:
        auth.set_custom_user_claims(target_uid, claims)
        print(f"[setUserClaims] uid={target_uid} claims={claims}")
        return https_fn.Response(json.dumps({"ok": True, "claims": claims}),
                                 status=200, content_type="application/json", headers=_CORS_CLAIMS)
    except Exception as e:
        print(f"[setUserClaims] ERROR: {e}")
        return err("Error interno del servidor", 500)


# ── CLOUD FUNCTION: MIGRAR CLAIMS EXISTENTES ─────────────────────────────────

@https_fn.on_request()
def migrarClaims(req: https_fn.Request) -> https_fn.Response:
    """
    Migración one-time: setea custom claims para todos los usuarios existentes.
    Solo puede ser invocado por supervisores activos.
    Recibe POST { idToken }
    """
    if req.method == "OPTIONS":
        return https_fn.Response("", status=204, headers=_CORS_CLAIMS)

    data = req.get_json(silent=True) or {}
    id_token = data.get("idToken", "")

    if not id_token:
        return https_fn.Response(json.dumps({"error": "idToken requerido"}), status=401,
                                 content_type="application/json", headers=_CORS_CLAIMS)
    try:
        decoded = auth.verify_id_token(id_token)
        caller_uid = decoded["uid"]
    except Exception:
        return https_fn.Response(json.dumps({"error": "Token inválido"}), status=401,
                                 content_type="application/json", headers=_CORS_CLAIMS)

    db = firestore.client()
    caller_doc = db.collection("usuarios").document(caller_uid).get()
    caller_data = caller_doc.to_dict() if caller_doc.exists else {}
    if caller_data.get("rol") != "supervisor" or not caller_data.get("activo"):
        return https_fn.Response(json.dumps({"error": "Solo supervisores activos"}), status=403,
                                 content_type="application/json", headers=_CORS_CLAIMS)

    usuarios = db.collection("usuarios").stream()
    migrados = []
    errores = []
    for doc in usuarios:
        d = doc.to_dict()
        uid = doc.id
        claims = {
            "rol": d.get("rol", "verificador"),
            "activo": bool(d.get("activo", True)),
        }
        try:
            auth.set_custom_user_claims(uid, claims)
            migrados.append({"uid": uid, "claims": claims})
        except Exception as e:
            errores.append({"uid": uid, "error": str(e)})

    return https_fn.Response(
        json.dumps({"ok": True, "migrados": len(migrados), "errores": errores}),
        status=200, content_type="application/json", headers=_CORS_CLAIMS
    )
